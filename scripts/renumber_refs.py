"""删除未引用的参考文献[2][12][15]，并按顺序重新编号全文引用和参考文献列表"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from lxml import etree
from docx.oxml.ns import qn

path = r'C:\Users\mc_leafwave\OneDrive\文档\统计建模\项目代码\paper\基于CHARLS的夫妻共病轨迹耦合及交互健康效应研究6(4).docx'
doc = Document(path)

# 删除的旧编号: 2, 12, 15
# 旧编号 -> 新编号 映射
old_to_new = {}
new_num = 1
for old_num in range(1, 28):
    if old_num in (2, 12, 15):
        old_to_new[old_num] = None  # 删除
    else:
        old_to_new[old_num] = new_num
        new_num += 1

print("编号映射 (旧->新):")
for old, new in old_to_new.items():
    if new is not None:
        print(f"  [{old}] -> [{new}]")
    else:
        print(f"  [{old}] -> 删除")

# ============================================================
# Step 1: 替换正文中所有引用标注的编号
# ============================================================
# 需要处理run级别的文本（因为引用是上标格式的独立run）

ref_para_start = None
for i, para in enumerate(doc.paragraphs):
    if '参考文献' in para.text.strip() and len(para.text.strip()) < 20:
        ref_para_start = i
        break

print(f"\n参考文献起始段落: [{ref_para_start}]")

# 处理正文中的引用标注 (段落0 到 ref_para_start)
def remap_citation(text):
    """将 [old1,old2] 格式的引用替换为 [new1,new2]"""
    def replace_match(m):
        content = m.group(1)
        nums = [int(n.strip()) for n in content.split(',')]
        new_nums = []
        for n in nums:
            mapped = old_to_new.get(n)
            if mapped is not None:
                new_nums.append(str(mapped))
        if new_nums:
            return '[' + ','.join(new_nums) + ']'
        else:
            return ''  # 引用被完全删除
    return re.sub(r'\[([\d,\s]+)\]', replace_match, text)

changes = 0
for i in range(ref_para_start):
    para = doc.paragraphs[i]
    for run in para.runs:
        old_text = run.text
        new_text = remap_citation(old_text)
        if old_text != new_text:
            run.text = new_text
            changes += 1

print(f"正文引用替换: {changes} 处")

# ============================================================
# Step 2: 删除参考文献列表中编号2, 12, 15的条目，并重新编号
# ============================================================
# 找到所有参考文献段落
ref_paras_to_delete = []
ref_paras_to_renumber = []

for i in range(ref_para_start + 1, len(doc.paragraphs)):
    para = doc.paragraphs[i]
    text = para.text.strip()
    if not text:
        continue
    # 匹配 [N] 或 N. 开头的参考文献
    m = re.match(r'^\[?(\d+)\]?\s*', text)
    if m:
        old_num = int(m.group(1))
        if old_num in (2, 12, 15):
            ref_paras_to_delete.append(i)
            print(f"  删除参考文献 [{old_num}] (段落 {i})")
        else:
            new_num = old_to_new.get(old_num)
            if new_num is not None and new_num != old_num:
                ref_paras_to_renumber.append((i, old_num, new_num))

# 删除参考文献段落 (从后往前删以免索引错位)
# python-docx 不支持直接删除段落，需要操作 XML
for i in sorted(ref_paras_to_delete, reverse=True):
    para = doc.paragraphs[i]
    p_element = para._element
    p_element.getparent().remove(p_element)
    print(f"  已删除段落 [{i}]")

# 重新编号剩余参考文献
# 需要重新读取段落（因为删除后索引变化）
# 直接在run中替换编号
ref_started = False
current_ref = 0
for para in doc.paragraphs:
    text = para.text.strip()
    if '参考文献' in text and len(text) < 20:
        ref_started = True
        continue
    if not ref_started:
        continue
    if not text:
        continue

    # 查找开头的 [N] 编号
    m = re.match(r'^\[?(\d+)\]?', text)
    if m:
        old_num = int(m.group(1))
        new_num_val = old_to_new.get(old_num)
        if new_num_val is not None:
            # 替换第一个run中的编号
            for run in para.runs:
                old_run_text = run.text
                # 替换开头的 [old] 为 [new]
                new_run_text = re.sub(
                    r'^\[?' + str(old_num) + r'\]?',
                    f'[{new_num_val}]',
                    old_run_text,
                    count=1
                )
                if new_run_text != old_run_text:
                    run.text = new_run_text
                    break

# ============================================================
# Step 3: 保存并验证
# ============================================================
doc.save(path)

# 验证
doc2 = Document(path)
all_nums = set()
ref_started2 = False
for para in doc2.paragraphs:
    text = para.text.strip()
    if '参考文献' in text and len(text) < 20:
        ref_started2 = True
        continue
    if not ref_started2:
        found = re.findall(r'\[[\d,]+\]', text)
        for f in found:
            for n in re.findall(r'\d+', f):
                all_nums.add(int(n))

# 统计参考文献条目数
ref_count = 0
for para in doc2.paragraphs:
    if ref_started2:
        if re.match(r'^\[?\d+\]', para.text.strip()):
            ref_count += 1
    if '参考文献' in para.text.strip() and len(para.text.strip()) < 20:
        ref_started2 = True

print(f"\n=== 验证 ===")
print(f"正文引用编号: {sorted(all_nums)}")
print(f"参考文献条目数: {ref_count}")
print(f"最大引用编号: {max(all_nums) if all_nums else 0}")
print("完成！")
