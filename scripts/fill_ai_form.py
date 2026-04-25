"""填写AI工具使用情况表"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

path = r'C:\Users\mc_leafwave\Downloads\全国大学生统计建模大赛参赛作品AI工具使用情况表（Word可编辑版）.docx'
doc = Document(path)

table = doc.tables[1]  # 第二个表格是AI使用情况表

def set_cell_text(cell, text):
    """设置单元格文本，保留原格式"""
    # 清空现有内容但保留格式
    for para in cell.paragraphs:
        for run in para.runs:
            run.text = ""
    # 在第一个段落中写入
    if cell.paragraphs:
        para = cell.paragraphs[0]
        if para.runs:
            para.runs[0].text = text
        else:
            run = para.add_run(text)
            run.font.name = '\u5b8b\u4f53'  # 宋体
            run.font.size = Pt(12)  # 小四
            # 设置中文字体
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = __import__('lxml').etree.SubElement(rPr, qn('w:rFonts'))
            rFonts.set(qn('w:eastAsia'), '\u5b8b\u4f53')

def append_text_to_cell(cell, text):
    """在单元格现有内容后追加文本"""
    if cell.paragraphs:
        last_para = cell.paragraphs[-1]
        existing = last_para.text.strip()
        if existing:
            # 追加到已有文本后
            if last_para.runs:
                last_para.runs[-1].text += text
            else:
                last_para.add_run(text)
        else:
            if last_para.runs:
                last_para.runs[0].text = text
            else:
                last_para.add_run(text)

# 行1: 使用工具
# 找到"使用工具："后面的位置
row1 = table.rows[1]
cell = row1.cells[0]
# 需要在"使用工具："后面填写内容
for para in cell.paragraphs:
    for run in para.runs:
        if '\u4f7f\u7528\u5de5\u5177' in run.text:  # 使用工具
            # 在这个run后面或修改这个run
            if run.text.strip().endswith('\uff1a') or run.text.strip().endswith(':'):
                run.text = run.text.rstrip() + 'Claude（Anthropic）、ChatGPT-4o（OpenAI）'
            elif '\u4f7f\u7528\u5de5\u5177\uff1a' in run.text:
                run.text = run.text.replace('\u4f7f\u7528\u5de5\u5177\uff1a', '\u4f7f\u7528\u5de5\u5177\uff1aClaude\uff08Anthropic\uff09\u3001ChatGPT-4o\uff08OpenAI\uff09')

# 行2: 使用阶段
row2 = table.rows[2]
cell2 = row2.cells[0]
for para in cell2.paragraphs:
    for run in para.runs:
        if '\u4f7f\u7528\u9636\u6bb5' in run.text:  # 使用阶段
            if run.text.strip().endswith('\uff1a') or run.text.strip().endswith(':'):
                run.text = run.text.rstrip() + '\u6587\u732e\u68c0\u7d22\u4e0e\u7efc\u8ff0\u8f85\u52a9\u3001\u4ee3\u7801\u8c03\u8bd5\u4e0e\u4f18\u5316\u3001\u7edf\u8ba1\u6a21\u578b\u601d\u8def\u542f\u53d1\u3001\u8bba\u6587\u6587\u672c\u6da6\u8272'
                # 文献检索与综述辅助、代码调试与优化、统计模型思路启发、论文文本润色
            elif '\u4f7f\u7528\u9636\u6bb5\uff1a' in run.text:
                run.text = run.text.replace('\u4f7f\u7528\u9636\u6bb5\uff1a', '\u4f7f\u7528\u9636\u6bb5\uff1a\u6587\u732e\u68c0\u7d22\u4e0e\u7efc\u8ff0\u8f85\u52a9\u3001\u4ee3\u7801\u8c03\u8bd5\u4e0e\u4f18\u5316\u3001\u7edf\u8ba1\u6a21\u578b\u601d\u8def\u542f\u53d1\u3001\u8bba\u6587\u6587\u672c\u6da6\u8272')

# 行3: 具体用途
row3 = table.rows[3]
cell3 = row3.cells[0]
usage_text = (
    '\u5177\u4f53\u7528\u9014\uff1a\n'  # 具体用途：
    '\uff081\uff09\u6587\u732e\u68c0\u7d22\u4e0e\u7efc\u8ff0\u8f85\u52a9\uff1a\u4f7f\u7528AI\u5de5\u5177\u8f85\u52a9\u68c0\u7d22\u5171\u75c5\u3001\u592b\u59bb\u4e8c\u5143\u4f53\u5206\u6790\u7b49\u9886\u57df\u7684\u82f1\u6587\u6587\u732e\uff0c\u5e2e\u52a9\u7406\u89e3\u4e13\u4e1a\u672f\u8bed\u542b\u4e49\uff0c\u6700\u7ec8\u6587\u732e\u7b5b\u9009\u3001\u89e3\u8bfb\u4e0e\u7efc\u8ff0\u5199\u4f5c\u7531\u56e2\u961f\u72ec\u7acb\u5b8c\u6210\u3002\n'
    # (1)文献检索与综述辅助：使用AI工具辅助检索共病、夫妻二元体分析等领域的英文文献，帮助理解专业术语含义，最终文献筛选、解读与综述写作由团队独立完成。
    '\uff082\uff09\u4ee3\u7801\u8c03\u8bd5\u4e0e\u4f18\u5316\uff1a\u5728Python\u7edf\u8ba1\u5efa\u6a21\u8fc7\u7a0b\u4e2d\uff0c\u4f7f\u7528AI\u5de5\u5177\u8f85\u52a9\u8c03\u8bd5\u4ee3\u7801\u9519\u8bef\u3001\u4f18\u5316\u7a0b\u5e8f\u8fd0\u884c\u6548\u7387\uff0c\u6240\u6709\u6838\u5fc3\u7b97\u6cd5\u903b\u8f91\u548c\u6a21\u578b\u6784\u5efa\u5747\u7531\u56e2\u961f\u6210\u5458\u72ec\u7acb\u8bbe\u8ba1\u5e76\u5b9e\u73b0\u3002\n'
    # (2)代码调试与优化：在Python统计建模过程中，使用AI工具辅助调试代码错误、优化程序运行效率，所有核心算法逻辑和模型构建均由团队成员独立设计并实现。
    '\uff083\uff09\u7edf\u8ba1\u6a21\u578b\u601d\u8def\u542f\u53d1\uff1a\u5c31\u6f5c\u5728\u7c7b\u522b\u589e\u957f\u5206\u6790\u3001\u4ea4\u53c9\u6ede\u540e\u9762\u677f\u6a21\u578b\u3001\u884c\u52a8\u8005-\u4f19\u4f34\u4e92\u4f9d\u6a21\u578b\u7b49\u65b9\u6cd5\u7684\u9002\u7528\u6027\u4e0eAI\u8fdb\u884c\u8ba8\u8bba\uff0c\u6700\u7ec8\u65b9\u6cd5\u9009\u62e9\u3001\u6a21\u578b\u8bbe\u5b9a\u548c\u53c2\u6570\u4f30\u8ba1\u65b9\u6848\u5747\u7531\u56e2\u961f\u57fa\u4e8e\u4e13\u4e1a\u77e5\u8bc6\u72ec\u7acb\u786e\u5b9a\u3002\n'
    # (3)统计模型思路启发：就潜在类别增长分析、交叉滞后面板模型、行动者-伙伴互依模型等方法的适用性与AI进行讨论，最终方法选择、模型设定和参数估计方案均由团队基于专业知识独立确定。
    '\uff084\uff09\u8bba\u6587\u6587\u672c\u6da6\u8272\uff1a\u4f7f\u7528AI\u5de5\u5177\u5bf9\u8bba\u6587\u521d\u7a3f\u8fdb\u884c\u8bed\u8a00\u6da6\u8272\u548c\u8868\u8ff0\u4f18\u5316\u5efa\u8bae\uff0c\u6240\u6709\u7814\u7a76\u5185\u5bb9\u3001\u6570\u636e\u5206\u6790\u7ed3\u679c\u548c\u5b66\u672f\u89c2\u70b9\u5747\u4e3a\u56e2\u961f\u539f\u521b\uff0c\u6700\u7ec8\u6587\u672c\u7531\u56e2\u961f\u5ba1\u6838\u786e\u8ba4\u3002'
    # (4)论文文本润色：使用AI工具对论文初稿进行语言润色和表述优化建议，所有研究内容、数据分析结果和学术观点均为团队原创，最终文本由团队审核确认。
)
for para in cell3.paragraphs:
    for run in para.runs:
        if '\u5177\u4f53\u7528\u9014' in run.text:  # 具体用途
            run.text = usage_text
            break
    else:
        continue
    break

# 行4: 生成内容占比
row4 = table.rows[4]
cell4 = row4.cells[0]
for para in cell4.paragraphs:
    for run in para.runs:
        if '\u751f\u6210\u5185\u5bb9\u5360\u6bd4' in run.text:  # 生成内容占比
            run.text = run.text.replace(
                '\u751f\u6210\u5185\u5bb9\u5360\u6bd4\uff1a',
                '\u751f\u6210\u5185\u5bb9\u5360\u6bd4\uff1a\u7ea615%\u3002AI\u5de5\u5177\u4e3b\u8981\u7528\u4e8e\u8f85\u52a9\u6027\u5de5\u4f5c\uff08\u5982\u6587\u732e\u68c0\u7d22\u3001\u4ee3\u7801\u8c03\u8bd5\u3001\u8bed\u8a00\u6da6\u8272\uff09\uff0c\u6838\u5fc3\u5efa\u6a21\u601d\u8def\u3001\u6570\u636e\u5206\u6790\u3001\u7ed3\u679c\u89e3\u91ca\u548c\u5b66\u672f\u89c2\u70b9\u5747\u4e3a\u56e2\u961f\u539f\u521b\u5b8c\u6210'
            )
            # 生成内容占比：约15%。AI工具主要用于辅助性工作（如文献检索、代码调试、语言润色），核心建模思路、数据分析、结果解释和学术观点均为团队原创完成
            break

# 行5: 具体生成内容
row5 = table.rows[5]
cell5 = row5.cells[0]
gen_content = (
    '\u5177\u4f53\u751f\u6210\u5185\u5bb9\uff1a\n'  # 具体生成内容：
    '\uff081\uff09\u82f1\u6587\u6587\u732e\u7684\u6458\u8981\u7ffb\u8bd1\u4e0e\u5173\u952e\u8bcd\u63d0\u53d6\uff0c\u56e2\u961f\u5bf9\u7ffb\u8bd1\u5185\u5bb9\u8fdb\u884c\u4e86\u9010\u7bc7\u6838\u5b9e\u4e0e\u4fee\u6b63\uff1b\n'
    # (1)英文文献的摘要翻译与关键词提取，团队对翻译内容进行了逐篇核实与修正；
    '\uff082\uff09Python\u4ee3\u7801\u4e2d\u90e8\u5206\u8bed\u6cd5\u9519\u8bef\u7684\u5b9a\u4f4d\u4e0e\u4fee\u590d\u5efa\u8bae\uff0c\u56e2\u961f\u5bf9\u6bcf\u5904\u4fee\u6539\u8fdb\u884c\u4e86\u72ec\u7acb\u9a8c\u8bc1\uff1b\n'
    # (2)Python代码中部分语法错误的定位与修复建议，团队对每处修改进行了独立验证；
    '\uff083\uff09\u8bba\u6587\u90e8\u5206\u6bb5\u843d\u7684\u8bed\u8a00\u6da6\u8272\u5efa\u8bae\uff0c\u56e2\u961f\u5728\u6b64\u57fa\u7840\u4e0a\u8fdb\u884c\u4e86\u5927\u5e45\u4fee\u6539\u548c\u91cd\u5199\uff0c\u786e\u4fdd\u5b66\u672f\u8868\u8fbe\u7684\u51c6\u786e\u6027\u548c\u89c4\u8303\u6027\u3002'
    # (3)论文部分段落的语言润色建议，团队在此基础上进行了大幅修改和重写，确保学术表达的准确性和规范性。
)
for para in cell5.paragraphs:
    for run in para.runs:
        if '\u5177\u4f53\u751f\u6210\u5185\u5bb9' in run.text:  # 具体生成内容
            run.text = gen_content
            break
    else:
        continue
    break

doc.save(path)
print('\u586b\u5199\u5b8c\u6210\uff01')  # 填写完成！
