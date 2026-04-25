"""删除3条未引用参考文献并重新编号 - 修复版"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

path = r'C:\Users\mc_leafwave\OneDrive\文档\统计建模\项目代码\paper\基于CHARLS的夫妻共病轨迹耦合及交互健康效应研究6(4).docx'
doc = Document(path)

old_to_new = {}
new_num = 1
for old_num in range(1, 28):
    if old_num in (2, 12, 15):
        old_to_new[old_num] = None
    else:
        old_to_new[old_num] = new_num
        new_num += 1

# 找参考文献起始
ref_start = None
for i, para in enumerate(doc.paragraphs):
    if '\u53c2\u8003\u6587\u732e' in para.text.strip() and len(para.text.strip()) < 20:
        ref_start = i
        break

# 先收集所有参考文献段落和操作
ref_paras = []
for i in range(ref_start + 1, len(doc.paragraphs)):
    para = doc.paragraphs[i]
    text = para.text.strip()
    if not text:
        continue
    # 检查是否到了附录
    if '\u9644\u5f55' in text:  # 附录
        break
    ref_paras.append((i, para))

print(f'\u53c2\u8003\u6587\u732e\u6761\u6570: {len(ref_paras)}')  # 参考文献条数

# 按旧编号处理
to_delete_elements = []
for idx, (para_i, para) in enumerate(ref_paras):
    old_num = idx + 1
    if old_num > 27:
        break

    if old_num in (2, 12, 15):
        to_delete_elements.append(para._element)
        print(f'  \u5220\u9664 [{old_num}]: {para.text[:50]}')  # 删除
    else:
        new = old_to_new[old_num]
        # 添加或替换编号
        if para.runs:
            first_run = para.runs[0]
            old_text = first_run.text
            # 去掉已有的 [N] 前缀
            cleaned = re.sub(r'^\[\d+\]\s*', '', old_text)
            first_run.text = f'[{new}] ' + cleaned

# 删除段落
for elem in to_delete_elements:
    elem.getparent().remove(elem)
print(f'\n\u5df2\u5220\u9664 {len(to_delete_elements)} \u6761')  # 已删除 N 条

doc.save(path)
print('\u4fdd\u5b58\u5b8c\u6210')  # 保存完成

# 验证
doc2 = Document(path)
ref_started = False
count = 0
for para in doc2.paragraphs:
    text = para.text.strip()
    if '\u53c2\u8003\u6587\u732e' in text and len(text) < 20:
        ref_started = True
        continue
    if '\u9644\u5f55' in text:
        break
    if ref_started and text:
        count += 1
        if count <= 3 or count >= 22:
            print(f'  [{count}] {text[:80]}')
        elif count == 4:
            print('  ...')

print(f'\n\u6700\u7ec8\u53c2\u8003\u6587\u732e: {count} \u6761')  # 最终参考文献: N 条

# 验证正文引用
all_nums = set()
for para in doc2.paragraphs:
    text = para.text.strip()
    if '\u53c2\u8003\u6587\u732e' in text and len(text) < 20:
        break
    for n in re.findall(r'\[(\d+)', text):
        all_nums.add(int(n))
print(f'\u6b63\u6587\u5f15\u7528\u7f16\u53f7: {sorted(all_nums)}')  # 正文引用编号
print(f'\u6700\u5927\u7f16\u53f7: {max(all_nums) if all_nums else 0}')  # 最大编号
